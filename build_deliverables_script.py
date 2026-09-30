import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from generate_all import PTA_CLIN_REAJUSTE_DATA

# Formatting helpers for docx
def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="D3D3D3", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = ['<w:tcBorders ' + nsdecls("w") + '>']
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            borders.append(f'<{border} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>')
        else:
            borders.append(f'<w:{side} w:val="none"/>')
    borders.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(borders)))

def make_table_robust(table):
    for row in table.rows:
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    header_tr = table.rows[0]._element.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

# ==========================================
# 1. PTA CLIN REAJUSTE (DOCX & XLSX)
# ==========================================
def build_pta_clin_reajuste_docx(filename):
    doc = Document()
    
    # Page setup landscape
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PLAN DE TRAVAIL ANNUEL CLIN / PCI 2026-2027 (RÉAJUSTÉ)\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    r_sub = p_title.add_run("CENTRE HOSPITALIER UNIVERSITAIRE DE LA MÈRE ET DE L'ENFANT LAGUNE (CHUMEL)\n"
                            "Alignement rigoureux des activités et des périodes sur le PTA BUDGÉTISÉ de référence")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    headers = [
        "Objectifs stratégiques",
        "Interventions stratégiques",
        "Activité / Sous-activité",
        "Responsabilité",
        "Délai d’exécution (Période)",
        "Indicateurs de base",
        "Indicateurs cible",
        "Source de financement"
    ]
    
    col_widths = [Inches(1.2), Inches(1.3), Inches(2.7), Inches(1.1), Inches(1.2), Inches(1.3), Inches(1.3), Inches(0.8)]
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        set_cell_borders(hdr_cells[i], color="0F203C", sz="6")
        hdr_cells[i].width = col_widths[i]
        
    prev_os = ""
    for idx, act in enumerate(PTA_CLIN_REAJUSTE_DATA):
        row_cells = table.add_row().cells
        
        # OS
        display_os = act["os"] if act["os"] != prev_os else ""
        prev_os = act["os"]
        
        data_vals = [
            display_os,
            act["inter"],
            act["act"],
            act["resp"],
            act["per"],
            act["ind_base"],
            act["ind_cible"],
            act["source"]
        ]
        
        bg_col = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        if display_os:
            bg_col = "F0F4F8"
            
        for i, val in enumerate(data_vals):
            row_cells[i].text = val
            p = row_cells[i].paragraphs[0]
            if i in [0, 3, 4, 7]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                if i == 0 and val:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
                elif i == 4:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(44, 62, 80)
            set_cell_background(row_cells[i], bg_col)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=90, right=90)
            set_cell_borders(row_cells[i], color="D3D3D3", sz="4")
            row_cells[i].width = col_widths[i]
            
    make_table_robust(table)
    doc.save(filename)
    print(f"Saved {filename}")

def build_pta_clin_reajuste_xlsx(filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PTA CLIN Réajusté"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "N°",
        "Objectifs stratégiques",
        "Interventions stratégiques",
        "Activité / Sous-activité",
        "Responsabilité",
        "Délai d’exécution (Période réajustée)",
        "Indicateurs de base",
        "Indicateurs cible",
        "Source de financement"
    ]
    
    # Title row
    ws.merge_cells("A1:I1")
    t_cell = ws["A1"]
    t_cell.value = "PLAN DE TRAVAIL ANNUEL CLIN / PCI CHUMEL 2026-2027 (RÉAJUSTÉ)"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    # Subtitle row
    ws.merge_cells("A2:I2")
    s_cell = ws["A2"]
    s_cell.value = "Matrice opérationnelle consolidée - Activités et périodes alignées sur le PTA BUDGÉTISÉ de référence"
    s_cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    s_cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    s_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    # Headers
    ws.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=Side(style="medium", color="0F203C"),
            bottom=Side(style="medium", color="0F203C"),
            left=Side(style="thin", color="0F203C"),
            right=Side(style="thin", color="0F203C")
        )
        
    thin_border = Border(
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3")
    )
    
    for idx, act in enumerate(PTA_CLIN_REAJUSTE_DATA, 1):
        row_num = idx + 3
        ws.row_dimensions[row_num].height = 36
        bg_color = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        
        vals = [
            idx,
            act["os"],
            act["inter"],
            act["act"],
            act["resp"],
            act["per"],
            act["ind_base"],
            act["ind_cible"],
            act["source"]
        ]
        
        for c_idx, val in enumerate(vals, 1):
            c = ws.cell(row=row_num, column=c_idx, value=val)
            c.font = Font(name="Calibri", size=9.5)
            c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            c.border = thin_border
            if c_idx in [1, 5, 6, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if c_idx == 6:
                c.font = Font(name="Calibri", size=9.5, bold=True, color="2C3E50")
                
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["G"].width = 32
    ws.column_dimensions["H"].width = 34
    ws.column_dimensions["I"].width = 16
    
    ws.freeze_panes = "A4"
    wb.save(filename)
    print(f"Saved {filename}")

build_pta_clin_reajuste_docx("PTA CLIN PCI CHUMEL 2026-2027 Réajusté.docx")
build_pta_clin_reajuste_xlsx("PTA CLIN PCI CHUMEL 2026-2027 Réajusté.xlsx")

# ==========================================
# 2. PLAN STRATEGIQUE (DOCX & XLSX)
# ==========================================
# Filter activities that are continuous or recurrent (rec == True)
STRATEGIC_DATA = [a for a in PTA_CLIN_REAJUSTE_DATA if a["rec"]]

# Map periodicity clean names
def get_strategic_periodicity(act):
    p = act["per"].lower()
    if "hebdo" in p:
        return "Hebdomadaire (Continue)"
    elif "mensuel" in p:
        return "Mensuelle (Continue)"
    elif "trimestriel" in p:
        return "Trimestrielle (Continue)"
    elif "continu" in p:
        return "Continue / Permanente"
    elif "5 mai" in p or "mai 2027" in p:
        return "Annuelle (Chaque 5 mai)"
    else:
        return "Annuelle (Récurrente chaque exercice)"

def build_plan_strategique_docx(filename):
    doc = Document()
    
    # Orientation landscape
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PLAN STRATÉGIQUE DE PRÉVENTION ET CONTRÔLE DES INFECTIONS (2026-2030)\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    r_sub = p_title.add_run("CENTRE HOSPITALIER UNIVERSITAIRE DE LA MÈRE ET DE L'ENFANT LAGUNE (CHUMEL)\n"
                            "Cadre Programmatique Pluriannuel fondé sur les Activités Continues et Récurrentes")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    # Intro box
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(6)
    p_intro.paragraph_format.space_after = Pt(12)
    r_intro = p_intro.add_run(
        "Note de cadrage stratégique : Le présent Plan Stratégique pérennise l'ensemble des interventions structurelles indispensables à la sécurité des soins au CHUMEL. "
        "Conformément aux orientations de la Direction et du CLIN, il retient exclusivement les activités continues, récurrentes et annuelles (gouvernance permanente, formation annuelle obligatoire, surveillance sentinelle hebdomadaire, audits réguliers, contrôle microbiologique de routine et maintenance WASH continue)."
    )
    r_intro.font.name = "Calibri"
    r_intro.font.size = Pt(9.5)
    r_intro.font.italic = True
    r_intro.font.color.rgb = RGBColor(44, 62, 80)
    
    headers = [
        "Axe Stratégique",
        "Intervention Clé",
        "Activité Continue / Récurrente",
        "Périodicité / Fréquence",
        "Responsable de pilotage",
        "Indicateurs de performance stratégique",
        "Financement"
    ]
    
    col_widths = [Inches(1.5), Inches(1.5), Inches(3.2), Inches(1.3), Inches(1.2), Inches(1.9), Inches(0.8)]
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        set_cell_borders(hdr_cells[i], color="0F203C", sz="6")
        hdr_cells[i].width = col_widths[i]
        
    prev_os = ""
    for idx, act in enumerate(STRATEGIC_DATA):
        row_cells = table.add_row().cells
        display_os = act["os"] if act["os"] != prev_os else ""
        prev_os = act["os"]
        
        freq = get_strategic_periodicity(act)
        
        data_vals = [
            display_os,
            act["inter"],
            act["act"],
            freq,
            act["resp"],
            f"Base: {act['ind_base']}\nCible: {act['ind_cible']}",
            act["source"]
        ]
        
        bg_col = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        if display_os:
            bg_col = "F0F4F8"
            
        for i, val in enumerate(data_vals):
            row_cells[i].text = val
            p = row_cells[i].paragraphs[0]
            if i in [0, 3, 4, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                if i == 0 and val:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
                elif i == 3:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(197, 160, 89) if "Annuelle" in val else RGBColor(44, 62, 80)
            set_cell_background(row_cells[i], bg_col)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=90, right=90)
            set_cell_borders(row_cells[i], color="D3D3D3", sz="4")
            row_cells[i].width = col_widths[i]
            
    make_table_robust(table)
    doc.save(filename)
    print(f"Saved {filename}")

def build_plan_strategique_xlsx(filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan Stratégique PCI"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "N°",
        "Axe Stratégique",
        "Intervention Clé",
        "Activité Continue / Récurrente",
        "Périodicité / Fréquence",
        "Responsable de pilotage",
        "Indicateurs de base",
        "Indicateurs cible",
        "Financement"
    ]
    
    # Title
    ws.merge_cells("A1:I1")
    t_cell = ws["A1"]
    t_cell.value = "PLAN STRATÉGIQUE DE PRÉVENTION ET CONTRÔLE DES INFECTIONS CHUMEL (2026-2030)"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells("A2:I2")
    s_cell = ws["A2"]
    s_cell.value = "Cadre programmatique pluriannuel des activités continues, permanentes et récurrentes annuelles"
    s_cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    s_cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    s_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    ws.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=Side(style="medium", color="0F203C"),
            bottom=Side(style="medium", color="0F203C"),
            left=Side(style="thin", color="0F203C"),
            right=Side(style="thin", color="0F203C")
        )
        
    thin_border = Border(
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3")
    )
    
    for idx, act in enumerate(STRATEGIC_DATA, 1):
        row_num = idx + 3
        ws.row_dimensions[row_num].height = 36
        bg_color = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        freq = get_strategic_periodicity(act)
        
        vals = [
            idx,
            act["os"],
            act["inter"],
            act["act"],
            freq,
            act["resp"],
            act["ind_base"],
            act["ind_cible"],
            act["source"]
        ]
        
        for c_idx, val in enumerate(vals, 1):
            c = ws.cell(row=row_num, column=c_idx, value=val)
            c.font = Font(name="Calibri", size=9.5)
            c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            c.border = thin_border
            if c_idx in [1, 5, 6, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if c_idx == 5:
                c.font = Font(name="Calibri", size=9.5, bold=True, color="1B365D")
                
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 30
    ws.column_dimensions["H"].width = 32
    ws.column_dimensions["I"].width = 16
    
    ws.freeze_panes = "A4"
    wb.save(filename)
    print(f"Saved {filename}")

build_plan_strategique_docx("PLAN STRATÉGIQUE PCI CHUMEL 2026-2030.docx")
build_plan_strategique_xlsx("PLAN STRATÉGIQUE PCI CHUMEL 2026-2030.xlsx")

# ==========================================
# 3. PLAN OPERATIONNEL (DOCX & XLSX)
# ==========================================
MONTHS = ["Oct", "Nov", "Déc", "Jan", "Fév", "Mar", "Avr", "Mai", "Jun", "Jul", "Aoû", "Sep"]

def get_months_active(per):
    p = per.lower()
    res = [False] * 12
    if "continue" in p or "continu" in p:
        return [True] * 12
    if "hebdo" in p or "mensuel" in p:
        return [True] * 12
    if "trimestriel" in p:
        # Nov/Dec, Feb/Mar, May/Jun, Aug/Sep
        res[1] = True # Nov
        res[2] = True # Dec
        res[4] = True # Fev
        res[5] = True # Mar
        res[7] = True # Mai
        res[8] = True # Jun
        res[10] = True # Aou
        res[11] = True # Sep
        return res
    if "septembre 2026" in p:
        # Before Oct, mark first month or special note
        res[0] = True
    if "octobre 2026" in p or "octorbre" in p:
        res[0] = True
    if "novembre 2026" in p:
        res[1] = True
    if "décembre 2026" in p or "decembre 2026" in p:
        res[2] = True
    if "janvier 2027" in p:
        res[3] = True
    if "février 2027" in p or "fevrier 2027" in p:
        res[4] = True
    if "mars 2027" in p:
        res[5] = True
    if "avril 2027" in p:
        res[6] = True
    if "mai 2027" in p:
        res[7] = True
    if "juin 2027" in p:
        res[8] = True
    if "juillet 2027" in p:
        res[9] = True
    if "août 2027" in p or "aout 2027" in p:
        res[10] = True
    if "septembre 2027" in p:
        res[11] = True
        
    # ranges
    if "décembre 2026 - janvier 2027" in p:
        res[2] = True
        res[3] = True
    if "janvier - mars 2027" in p:
        res[3] = True
        res[4] = True
        res[5] = True
    if "février - mars 2027" in p or "février – mars 2027" in p:
        res[4] = True
        res[5] = True
    if "mars - avril 2027" in p or "mars – avril 2027" in p:
        res[5] = True
        res[6] = True
    if "février - avril 2027" in p:
        res[4] = True
        res[5] = True
        res[6] = True
    if "mai - juin 2027" in p:
        res[7] = True
        res[8] = True
    if "janvier - août 2027" in p:
        for m in range(3, 11):
            res[m] = True
    if "novembre - décembre 2026" in p:
        res[1] = True
        res[2] = True
    if "décembre 2026 - mars 2027" in p:
        for m in range(2, 6):
            res[m] = True
    return res

def build_plan_operationnel_docx(filename):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)
        section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PLAN OPÉRATIONNEL CLIN / PCI CHUMEL 2026-2027\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    r_sub = p_title.add_run("Planification détaillée d'exécution avec Périodes Exactes et Chronogramme Mensuel d'Exécution\n"
                            "Période couverte : Octobre 2026 – Septembre 2027")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    headers = [
        "N°",
        "Activité Opérationnelle",
        "Responsable",
        "Période exacte",
        "O", "N", "D", "J", "F", "M", "A", "M", "J", "J", "A", "S",
        "Livrable Opérationnel Attendu"
    ]
    
    col_widths = [
        Inches(0.4),  # N
        Inches(3.2),  # Act
        Inches(1.3),  # Resp
        Inches(1.4),  # Periode
        # 12 months (0.24 each = 2.88)
        Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24),
        Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24), Inches(0.24),
        Inches(1.8)   # Livrable
    ]
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=40, right=40)
        set_cell_borders(hdr_cells[i], color="0F203C", sz="4")
        hdr_cells[i].width = col_widths[i]
        
    for idx, act in enumerate(PTA_CLIN_REAJUSTE_DATA, 1):
        row_cells = table.add_row().cells
        m_active = get_months_active(act["per"])
        
        # Base texts
        row_cells[0].text = str(idx)
        row_cells[1].text = act["act"]
        row_cells[2].text = act["resp"]
        row_cells[3].text = act["per"]
        
        # 12 months
        for m_idx in range(12):
            cell = row_cells[4 + m_idx]
            if m_active[m_idx]:
                cell.text = "●"
                set_cell_background(cell, "C5A059") # Gold dot
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(8)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                cell.text = ""
                bg = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=20, right=20)
            set_cell_borders(cell, color="D3D3D3", sz="4")
            cell.width = col_widths[4 + m_idx]
            
        row_cells[16].text = act["ind_cible"]
        
        bg_col = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        for i in [0, 1, 2, 3, 16]:
            c = row_cells[i]
            p = c.paragraphs[0]
            if i in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8)
                if i == 3:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=60, bottom=60, left=60, right=60)
            set_cell_borders(c, color="D3D3D3", sz="4")
            c.width = col_widths[i]
            
    make_table_robust(table)
    doc.save(filename)
    print(f"Saved {filename}")

def build_plan_operationnel_xlsx(filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan Opérationnel 2026-2027"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        "N°",
        "Axe Stratégique",
        "Activité Opérationnelle",
        "Responsable",
        "Période d'exécution exacte",
        "Oct 26", "Nov 26", "Déc 26", "Jan 27", "Fév 27", "Mar 27",
        "Avr 27", "Mai 27", "Juin 27", "Juil 27", "Août 27", "Sept 27",
        "Livrable Opérationnel Attendu",
        "Source de Financement"
    ]
    
    # Title
    ws.merge_cells("A1:S1")
    t_cell = ws["A1"]
    t_cell.value = "PLAN OPÉRATIONNEL CLIN / PCI CHUMEL 2026-2027"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells("A2:S2")
    s_cell = ws["A2"]
    s_cell.value = "Planification chronologique mensuelle d'exécution des activités (Octobre 2026 – Septembre 2027)"
    s_cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    s_cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    s_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    ws.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        fill_col = "C5A059" if 6 <= col_idx <= 17 else "1B365D"
        cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=Side(style="medium", color="0F203C"),
            bottom=Side(style="medium", color="0F203C"),
            left=Side(style="thin", color="0F203C"),
            right=Side(style="thin", color="0F203C")
        )
        
    thin_border = Border(
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3")
    )
    
    for idx, act in enumerate(PTA_CLIN_REAJUSTE_DATA, 1):
        row_num = idx + 3
        ws.row_dimensions[row_num].height = 34
        bg_color = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        m_active = get_months_active(act["per"])
        
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=2, value=act["os"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row_num, column=3, value=act["act"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row_num, column=4, value=act["resp"]).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        c_per = ws.cell(row=row_num, column=5, value=act["per"])
        c_per.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c_per.font = Font(name="Calibri", size=9.5, bold=True, color="1B365D")
        
        for m_idx in range(12):
            c_m = ws.cell(row=row_num, column=6 + m_idx)
            c_m.border = thin_border
            if m_active[m_idx]:
                c_m.value = "X"
                c_m.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
                c_m.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
                c_m.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c_m.value = ""
                c_m.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                
        ws.cell(row=row_num, column=18, value=act["ind_cible"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row_num, column=19, value=act["source"]).alignment = Alignment(horizontal="center", vertical="center")
        
        for c_idx in [1, 2, 3, 4, 5, 18, 19]:
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = Font(name="Calibri", size=9) if c_idx != 5 else Font(name="Calibri", size=9.5, bold=True, color="1B365D")
            cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            cell.border = thin_border
            
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 46
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 24
    for col_l in ["F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]:
        ws.column_dimensions[col_l].width = 8
    ws.column_dimensions["R"].width = 36
    ws.column_dimensions["S"].width = 16
    
    ws.freeze_panes = "F4"
    wb.save(filename)
    print(f"Saved {filename}")

build_plan_operationnel_docx("PLAN OPÉRATIONNEL PCI CHUMEL 2026-2027.docx")
build_plan_operationnel_xlsx("PLAN OPÉRATIONNEL PCI CHUMEL 2026-2027.xlsx")

# ==========================================
# 4. PTA BUDGETISE FINAL (DOCX & XLSX)
# ==========================================
# Filter strictly activities with bud > 0
BUDGETED_ACTIVITIES = [a for a in PTA_CLIN_REAJUSTE_DATA if a["bud"] > 0]

def build_pta_budgetise_final_docx(filename):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PTA BUDGÉTISÉ FINAL CLIN / PCI CHUMEL 2026-2027\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)
    
    r_sub = p_title.add_run("CENTRE HOSPITALIER UNIVERSITAIRE DE LA MÈRE ET DE L'ENFANT LAGUNE (CHUMEL)\n"
                            "Plan d'Action Financier Consolidé – Uniquement les Activités Effectivement Budgétisées")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    # Summary note
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(10)
    r_n = p_note.add_run(
        "Champ d'application : Ce document constitue la version finale du Plan Budgétisé PCI 2026-2027. "
        "Conformément aux directives de cadrage budgétaire, il exclut les activités sans dotation financière directe et rassemble exclusivement les 17 activités ayant fait l'objet d'un chiffrage budgétaire validé, pour un montant total consolidé de 13 490 000 FCFA."
    )
    r_n.font.name = "Calibri"
    r_n.font.size = Pt(9.5)
    r_n.font.italic = True
    r_n.font.color.rgb = RGBColor(44, 62, 80)
    
    headers = [
        "N°",
        "Objectifs stratégiques",
        "Activité / Sous-activité budgétisée",
        "Responsabilité",
        "Période d’exécution",
        "Détail du calcul / Clé de coût",
        "Budget validé (FCFA)",
        "Source"
    ]
    
    col_widths = [Inches(0.4), Inches(1.8), Inches(3.2), Inches(1.2), Inches(1.4), Inches(1.8), Inches(1.1), Inches(0.8)]
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=80, right=80)
        set_cell_borders(hdr_cells[i], color="0F203C", sz="6")
        hdr_cells[i].width = col_widths[i]
        
    total_budget = 0
    prev_os = ""
    for idx, act in enumerate(BUDGETED_ACTIVITIES, 1):
        row_cells = table.add_row().cells
        display_os = act["os"] if act["os"] != prev_os else ""
        prev_os = act["os"]
        
        total_budget += act["bud"]
        
        data_vals = [
            str(idx),
            display_os,
            act["act"],
            act["resp"],
            act["per"],
            act["bud_detail"],
            f"{act['bud']:,d}".replace(",", " "),
            act["source"]
        ]
        
        bg_col = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        if display_os:
            bg_col = "F0F4F8"
            
        for i, val in enumerate(data_vals):
            row_cells[i].text = val
            p = row_cells[i].paragraphs[0]
            if i in [0, 3, 4, 7]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i == 6:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                if i == 1 and val:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
                elif i == 6:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
            set_cell_background(row_cells[i], bg_col)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=80, right=80)
            set_cell_borders(row_cells[i], color="D3D3D3", sz="4")
            row_cells[i].width = col_widths[i]
            
    # Total row
    tot_cells = table.add_row().cells
    tot_cells[0].text = ""
    tot_cells[1].text = "TOTAL GÉNÉRAL CONSOLIDÉ"
    tot_cells[2].text = "Ensemble des 17 activités budgétisées"
    tot_cells[3].text = ""
    tot_cells[4].text = ""
    tot_cells[5].text = ""
    tot_cells[6].text = f"{total_budget:,d}".replace(",", " ") + " FCFA"
    tot_cells[7].text = "FP / PTF"
    
    for i, c in enumerate(tot_cells):
        set_cell_background(c, "C5A059") # Gold
        set_cell_margins(c, top=100, bottom=100, left=80, right=80)
        set_cell_borders(c, color="A07D34", sz="6")
        p = c.paragraphs[0]
        if i in [1, 2]:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif i == 6:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        c.width = col_widths[i]
        
    make_table_robust(table)
    
    # Financial synthesis by Objective
    p_syn = doc.add_paragraph()
    p_syn.paragraph_format.space_before = Pt(16)
    p_syn.paragraph_format.space_after = Pt(6)
    r_syn = p_syn.add_run("SYNTHÈSE BUDGÉTAIRE PAR OBJECTIF STRATÉGIQUE")
    r_syn.font.name = "Calibri"
    r_syn.font.size = Pt(12)
    r_syn.font.bold = True
    r_syn.font.color.rgb = RGBColor(27, 54, 93)
    
    syn_headers = ["Objectif Stratégique", "Nombre d'activités budgétisées", "Montant (FCFA)", "Part (%)"]
    syn_widths = [Inches(4.5), Inches(2.2), Inches(2.2), Inches(1.5)]
    
    t_syn = doc.add_table(rows=1, cols=len(syn_headers))
    t_syn.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_syn.autofit = False
    
    hdr_syn = t_syn.rows[0].cells
    for i, h in enumerate(syn_headers):
        hdr_syn[i].text = h
        p = hdr_syn[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_syn[i], "2C3E50")
        set_cell_margins(hdr_syn[i], top=100, bottom=100, left=80, right=80)
        set_cell_borders(hdr_syn[i], color="1B2838", sz="6")
        hdr_syn[i].width = syn_widths[i]
        
    # Group by OS
    os_groups = {}
    for a in BUDGETED_ACTIVITIES:
        os_name = a["os"]
        if os_name not in os_groups:
            os_groups[os_name] = {"count": 0, "total": 0}
        os_groups[os_name]["count"] += 1
        os_groups[os_name]["total"] += a["bud"]
        
    for idx, (os_name, data) in enumerate(os_groups.items()):
        row = t_syn.add_row().cells
        pct = (data["total"] / total_budget) * 100
        vals = [
            os_name,
            str(data["count"]),
            f"{data['total']:,d}".replace(",", " "),
            f"{pct:.1f} %"
        ]
        bg = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(vals):
            row[i].text = val
            p = row[i].paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                if i in [2, 3]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(27, 54, 93)
            set_cell_background(row[i], bg)
            set_cell_margins(row[i], top=70, bottom=70, left=80, right=80)
            set_cell_borders(row[i], color="D3D3D3", sz="4")
            row[i].width = syn_widths[i]
            
    # Synthesis total row
    tot_syn_cells = t_syn.add_row().cells
    tot_syn_cells[0].text = "TOTAL GÉNÉRAL"
    tot_syn_cells[1].text = str(len(BUDGETED_ACTIVITIES))
    tot_syn_cells[2].text = f"{total_budget:,d}".replace(",", " ") + " FCFA"
    tot_syn_cells[3].text = "100.0 %"
    
    for i, c in enumerate(tot_syn_cells):
        set_cell_background(c, "1B365D")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        set_cell_borders(c, color="0F203C", sz="6")
        p = c.paragraphs[0]
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif i == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        c.width = syn_widths[i]
        
    make_table_robust(t_syn)
    doc.save(filename)
    print(f"Saved {filename}")

def build_pta_budgetise_final_xlsx(filename):
    wb = openpyxl.Workbook()
    
    # Sheet 1: Activités Budgétisées
    ws1 = wb.active
    ws1.title = "PTA Budgétisé (Activités)"
    ws1.views.sheetView[0].showGridLines = True
    
    headers = [
        "N°",
        "Objectif Stratégique",
        "Activité / Sous-activité budgétisée",
        "Responsable",
        "Période d'exécution",
        "Détail du calcul / Clé de coût",
        "Budget Alloué (FCFA)",
        "Source de Financement"
    ]
    
    # Title
    ws1.merge_cells("A1:H1")
    t1 = ws1["A1"]
    t1.value = "PTA BUDGÉTISÉ FINAL CLIN / PCI CHUMEL 2026-2027"
    t1.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t1.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    t1.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 30
    
    ws1.merge_cells("A2:H2")
    s1 = ws1["A2"]
    s1.value = "Plan financier définitif – Uniquement les 17 activités ayant fait l'objet d'un budget validé"
    s1.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    s1.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    s1.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20
    
    ws1.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=Side(style="medium", color="0F203C"),
            bottom=Side(style="medium", color="0F203C"),
            left=Side(style="thin", color="0F203C"),
            right=Side(style="thin", color="0F203C")
        )
        
    thin_border = Border(
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3")
    )
    
    for idx, act in enumerate(BUDGETED_ACTIVITIES, 1):
        row_num = idx + 3
        ws1.row_dimensions[row_num].height = 32
        bg_color = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        
        ws1.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=2, value=act["os"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=row_num, column=3, value=act["act"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=row_num, column=4, value=act["resp"]).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        c_per = ws1.cell(row=row_num, column=5, value=act["per"])
        c_per.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c_per.font = Font(name="Calibri", size=9.5, bold=True, color="1B365D")
        
        ws1.cell(row=row_num, column=6, value=act["bud_detail"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        c_bud = ws1.cell(row=row_num, column=7, value=act["bud"])
        c_bud.alignment = Alignment(horizontal="right", vertical="center")
        c_bud.font = Font(name="Calibri", size=9.5, bold=True, color="1B365D")
        c_bud.number_format = '#,##0 "FCFA"'
        
        ws1.cell(row=row_num, column=8, value=act["source"]).alignment = Alignment(horizontal="center", vertical="center")
        
        for c_idx in range(1, 9):
            c = ws1.cell(row=row_num, column=c_idx)
            c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            c.border = thin_border
            if c_idx not in [5, 7]:
                c.font = Font(name="Calibri", size=9)
                
    # Total row
    tot_row = len(BUDGETED_ACTIVITIES) + 4
    ws1.row_dimensions[tot_row].height = 30
    ws1.merge_cells(f"A{tot_row}:F{tot_row}")
    t_label = ws1.cell(row=tot_row, column=1, value="TOTAL GÉNÉRAL CONSOLIDÉ (17 Activités Budgétisées)")
    t_label.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    t_label.fill = PatternFill(start_color="C5A059", end_color="C5A059", fill_type="solid")
    t_label.alignment = Alignment(horizontal="right", vertical="center")
    
    t_val = ws1.cell(row=tot_row, column=7, value=f"=SUM(G4:G{tot_row-1})")
    t_val.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    t_val.fill = PatternFill(start_color="C5A059", end_color="C5A059", fill_type="solid")
    t_val.alignment = Alignment(horizontal="right", vertical="center")
    t_val.number_format = '#,##0 "FCFA"'
    
    t_src = ws1.cell(row=tot_row, column=8, value="FP / PTF")
    t_src.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    t_src.fill = PatternFill(start_color="C5A059", end_color="C5A059", fill_type="solid")
    t_src.alignment = Alignment(horizontal="center", vertical="center")
    
    medium_border = Border(
        top=Side(style="medium", color="A07D34"),
        bottom=Side(style="medium", color="A07D34"),
        left=Side(style="thin", color="A07D34"),
        right=Side(style="thin", color="A07D34")
    )
    for c_idx in range(1, 9):
        ws1.cell(row=tot_row, column=c_idx).border = medium_border
        
    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 24
    ws1.column_dimensions["C"].width = 46
    ws1.column_dimensions["D"].width = 20
    ws1.column_dimensions["E"].width = 24
    ws1.column_dimensions["F"].width = 38
    ws1.column_dimensions["G"].width = 22
    ws1.column_dimensions["H"].width = 16
    
    ws1.freeze_panes = "A4"
    
    # Sheet 2: Synthèse Budgétaire
    ws2 = wb.create_sheet(title="Synthèse Budgétaire")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:E1")
    ws2["A1"].value = "SYNTHÈSE BUDGÉTAIRE DU PTA PCI CHUMEL 2026-2027"
    ws2["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws2["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30
    
    syn_headers = ["N°", "Objectif Stratégique", "Nombre d'activités", "Budget Alloué (FCFA)", "Part (%)"]
    ws2.row_dimensions[3].height = 26
    for col_idx, h in enumerate(syn_headers, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            top=Side(style="medium", color="1B2838"),
            bottom=Side(style="medium", color="1B2838"),
            left=Side(style="thin", color="1B2838"),
            right=Side(style="thin", color="1B2838")
        )
        
    # Group OS
    os_dict = {}
    for a in BUDGETED_ACTIVITIES:
        o = a["os"]
        if o not in os_dict:
            os_dict[o] = {"count": 0, "sum": 0}
        os_dict[o]["count"] += 1
        os_dict[o]["sum"] += a["bud"]
        
    total_val = sum(a["bud"] for a in BUDGETED_ACTIVITIES)
    for idx, (o_name, o_data) in enumerate(os_dict.items(), 1):
        r_num = idx + 3
        ws2.row_dimensions[r_num].height = 26
        bg_col = "F9FBFC" if idx % 2 == 1 else "FFFFFF"
        
        ws2.cell(row=r_num, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r_num, column=2, value=o_name).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(row=r_num, column=3, value=o_data["count"]).alignment = Alignment(horizontal="center", vertical="center")
        
        c_amt = ws2.cell(row=r_num, column=4, value=o_data["sum"])
        c_amt.alignment = Alignment(horizontal="right", vertical="center")
        c_amt.font = Font(name="Calibri", size=10, bold=True, color="1B365D")
        c_amt.number_format = '#,##0 "FCFA"'
        
        c_pct = ws2.cell(row=r_num, column=5, value=f"=D{r_num}/$D${len(os_dict)+4}")
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.font = Font(name="Calibri", size=10, bold=True)
        c_pct.number_format = "0.0 %"
        
        for c_idx in range(1, 6):
            c = ws2.cell(row=r_num, column=c_idx)
            c.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
            c.border = thin_border
            if c_idx not in [4, 5]:
                c.font = Font(name="Calibri", size=9.5)
                
    # Total row
    syn_tot_row = len(os_dict) + 4
    ws2.row_dimensions[syn_tot_row].height = 30
    ws2.merge_cells(f"A{syn_tot_row}:B{syn_tot_row}")
    ws2.cell(row=syn_tot_row, column=1, value="TOTAL GÉNÉRAL CONSOLIDÉ").alignment = Alignment(horizontal="right", vertical="center")
    ws2.cell(row=syn_tot_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws2.cell(row=syn_tot_row, column=1).fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    
    ws2.cell(row=syn_tot_row, column=3, value=f"=SUM(C4:C{syn_tot_row-1})").alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=syn_tot_row, column=3).font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws2.cell(row=syn_tot_row, column=3).fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    
    c_tot_amt = ws2.cell(row=syn_tot_row, column=4, value=f"=SUM(D4:D{syn_tot_row-1})")
    c_tot_amt.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_amt.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c_tot_amt.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    c_tot_amt.number_format = '#,##0 "FCFA"'
    
    c_tot_pct = ws2.cell(row=syn_tot_row, column=5, value=1.0)
    c_tot_pct.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_pct.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    c_tot_pct.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    c_tot_pct.number_format = "0.0 %"
    
    for c_idx in range(1, 6):
        ws2.cell(row=syn_tot_row, column=c_idx).border = Border(
            top=Side(style="medium", color="0F203C"),
            bottom=Side(style="medium", color="0F203C"),
            left=Side(style="thin", color="0F203C"),
            right=Side(style="thin", color="0F203C")
        )
        
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 38
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 24
    ws2.column_dimensions["E"].width = 16
    
    wb.save(filename)
    print(f"Saved {filename}")

build_pta_budgetise_final_docx("PTA BUDGÉTISÉ FINAL CLIN PCI CHUMEL 2026-2027.docx")
build_pta_budgetise_final_xlsx("PTA BUDGÉTISÉ FINAL CLIN PCI CHUMEL 2026-2027.xlsx")

# Update repository files in place as well
build_pta_clin_reajuste_docx("PTA CLIN PCI CHUMEL 2026-2027 finale.docx")
build_pta_budgetise_final_docx("PTA BUDGETISER PCI CHUMEL 2026-2027 finale.docx")
